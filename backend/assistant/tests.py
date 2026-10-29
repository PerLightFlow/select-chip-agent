import io
import json
import tempfile
from datetime import date
from unittest.mock import patch, Mock
from django.test import SimpleTestCase, TransactionTestCase, override_settings
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import IntegrityError, transaction
from rest_framework.test import APIClient
from assistant.models import Conversation, Message, Turn, EnterpriseProject, ProjectVersion, Proposal, KnowledgeDocument, KnowledgeVersion
from assistant.knowledge import propose, confirm_proposal, project_record, retrieve, index_document, confirm_document
from assistant.runtime import run_turn
from assistant.tools import search_chips, public_url, execute
from assistant.management.commands.agent_worker import claim


class PublicEvidenceTests(SimpleTestCase):
    def setUp(self):
        from assistant import tools
        tools._research_cache.clear()

    def test_other_chip_datasheet_cannot_replace_requested_model(self):
        from assistant.tools import _research_query
        wrong = {'url':'https://documentation.espressif.com/esp32_datasheet_en.pdf', 'title':'ESP32 Datasheet', 'text':'ESP32', 'kind':'search_excerpt'}
        right = {'url':'https://documentation.espressif.com/esp32-p4_datasheet_en.pdf', 'title':'ESP32-P4 Datasheet', 'text':'ESP32-P4', 'kind':'search_excerpt'}
        evidence = {'url':right['url'], 'text':'ESP32-P4 official text '*20, 'notice':'', 'kind':'source_text'}
        with patch('assistant.tools._search_query',return_value={'items':[wrong,right]}), patch('assistant.tools.read_source',return_value=evidence) as read:
            result = _research_query('ESP32-P4')
        read.assert_called_once_with(right['url'])
        self.assertEqual(result['items'],[right])
        self.assertEqual(result['evidence'][0]['url'],right['url'])

    def test_irrelevant_search_results_are_not_evidence(self):
        from assistant.tools import _research_query
        wrong = {'url':'https://example.com/people', 'title':'People directory', 'text':'unrelated', 'kind':'search_excerpt'}
        with patch('assistant.tools._search_query',return_value={'items':[wrong]}), patch('assistant.tools.read_source') as read:
            result = _research_query('RK3568')
        read.assert_not_called()
        self.assertEqual(result['evidence'],[])
        self.assertTrue(result['warnings'])


class AppServerBudgetTests(SimpleTestCase):
    def test_exhausted_tool_budget_still_allows_grounded_final_answer(self):
        from assistant.app_server import run_codex
        class Output(io.StringIO):
            def close(self):
                self.recorded = self.getvalue()
                super().close()
        stream = [{'id':1,'result':{}}, {'id':2,'result':{'thread':{'id':'thread'}}}, {'id':3,'result':{}}]
        stream += [{'id':100+i,'method':'item/tool/call','params':{'tool':'lookup','arguments':{'n':i}}} for i in range(3)]
        stream += [{'method':'thread/tokenUsage/updated','params':{'tokenUsage':{'total':{'inputTokens':123,'outputTokens':45}}}},
                   {'method':'item/completed','params':{'item':{'type':'agentMessage','text':'已核实 A 的接口，B 和 C 尚缺证据。'}}},
                   {'method':'turn/completed','params':{'turn':{'status':'completed'}}}]
        proc=Mock();proc.stdin=Output();proc.stdout=io.StringIO('\n'.join(json.dumps(x) for x in stream)+'\n');proc.poll.return_value=0
        execute=Mock(return_value={'source':'official fixture'})
        with patch('assistant.app_server.subprocess.Popen',return_value=proc):
            text,usage=run_codex(binary='fixture',base_url='https://example.test',api_key='test-only',model='fixture',instructions='test',history=[],
                tools=[{'function':{'name':'lookup','description':'test','parameters':{'type':'object'}}}],
                execute_tool=execute,progress=lambda _:None,check_active=lambda:None,max_tools=1)
        self.assertEqual(execute.call_count,1)
        self.assertIn('尚缺证据',text);self.assertEqual(usage['input_tokens'],123)
        replies=[json.loads(line) for line in proc.stdin.recorded.splitlines() if '"contentItems"' in line]
        self.assertEqual([x['result']['success'] for x in replies],[True,False,False])
        self.assertIn('预算已用完',replies[-1]['result']['contentItems'][0]['text'])


def card(code='PROJ-A', chip='ESP32-P4', effective='2026-01-01'):
    return {'code':code,'name':'项目 A','aliases':['视觉设备'],'summary':'企业测试项目',
        'modules':[{'name':'图像采集','chips':[chip],'responsible':[{'name':'张工','department':'硬件组','work':'电路设计'}],
                    'participants':[{'name':'李工','department':'软件组','work':'驱动'}],'notes':''}],
        'document_ids':[],'evidence_note':'用户明确提供','effective_from':effective}


@override_settings(AGENT_RUNTIME='chat', REST_FRAMEWORK={'DEFAULT_AUTHENTICATION_CLASSES':['rest_framework.authentication.SessionAuthentication'], 'DEFAULT_PERMISSION_CLASSES':['rest_framework.permissions.IsAuthenticated']})
class WorkspaceTests(TransactionTestCase):
    def setUp(self):
        self.alice=get_user_model().objects.create_user('alice',password='example-password')
        self.bob=get_user_model().objects.create_user('bob',password='example-password')
        self.admin=get_user_model().objects.create_user('admin',is_staff=True)
        self.client=APIClient();self.client.force_authenticate(self.alice)
        self.media=tempfile.TemporaryDirectory()
        self.override=override_settings(MEDIA_ROOT=self.media.name);self.override.enable()
        self.addCleanup(self.override.disable);self.addCleanup(self.media.cleanup)

    def create_project(self):
        return confirm_proposal(propose(self.alice,card()).pk,self.alice).project

    def document(self):
        d=KnowledgeDocument.objects.create(owner=self.alice,title='项目 A 设计')
        v=KnowledgeVersion.objects.create(document=d,revision=1,text='项目 A 图像采集采用 ESP32-P4，硬件组张工负责电路设计。',state='indexing')
        index_document(v.pk);v.refresh_from_db()
        return d,v

    def test_private_conversations_including_admin(self):
        c=Conversation.objects.create(owner=self.alice)
        Message.objects.create(conversation=c,role='user',content='私有探索')
        for user in [self.bob,self.admin]:
            self.client.force_authenticate(user)
            self.assertEqual(self.client.get(f'/api/assistant/conversations/{c.pk}/').status_code,404)
            self.assertEqual(self.client.post(f'/api/assistant/conversations/{c.pk}/messages/',{'text':'越权'},format='json').status_code,404)
            self.assertEqual(self.client.get('/api/assistant/conversations/').data,[])

    def test_authentication_required(self):
        self.client.force_authenticate(None)
        for path in ['conversations','documents','projects','proposals']:
            self.assertEqual(self.client.get(f'/api/assistant/{path}/').status_code,403)

    def test_one_active_turn_and_input_preserved(self):
        c=Conversation.objects.create(owner=self.alice)
        endpoint=f'/api/assistant/conversations/{c.pk}/messages/'
        self.assertEqual(self.client.post(endpoint,{'text':'想做视觉设备'},format='json').status_code,202)
        self.assertEqual(self.client.post(endpoint,{'text':'另一个请求'},format='json').status_code,409)
        self.assertEqual(c.messages.count(),1)
        with self.assertRaises(IntegrityError), transaction.atomic():
            Turn.objects.create(conversation=c)

    def test_worker_claim_and_cancel(self):
        c=Conversation.objects.create(owner=self.alice)
        t=Turn.objects.create(conversation=c)
        job=claim();self.assertEqual(job[1],t.pk);self.assertIsNone(claim())
        self.client.force_authenticate(self.bob)
        self.assertEqual(self.client.post(f'/api/assistant/turns/{t.pk}/stop/',{},format='json').status_code,404)
        self.client.force_authenticate(self.alice)
        self.assertEqual(self.client.post(f'/api/assistant/turns/{t.pk}/stop/',{},format='json').data['state'],'cancelled')
        with patch('assistant.runtime.chat') as chat:
            run_turn(t.pk)
        chat.assert_not_called();self.assertEqual(c.messages.count(),0)

    def test_cancel_during_model_does_not_write_answer(self):
        c=Conversation.objects.create(owner=self.alice);t=Turn.objects.create(conversation=c,state='running')
        def finish(*a,**k):
            Turn.objects.filter(pk=t.pk).update(state='cancelled')
            return {'role':'assistant','content':'不能保存这个结果'},{}
        with patch('assistant.runtime.chat',side_effect=finish):run_turn(t.pk)
        self.assertFalse(c.messages.exists())

    def test_multiturn_question_and_tool_query(self):
        self.create_project()
        c=Conversation.objects.create(owner=self.alice)
        Message.objects.create(conversation=c,role='user',content='帮我找项目芯片')
        t=Turn.objects.create(conversation=c,state='running')
        with patch('assistant.runtime.chat',return_value=({'role':'assistant','content':'你想查哪个项目？'},{})):run_turn(t.pk)
        Message.objects.create(conversation=c,role='user',content='项目 A')
        t2=Turn.objects.create(conversation=c,state='running')
        tool={'role':'assistant','content':None,'tool_calls':[{'id':'call_1','type':'function','function':{'name':'query_projects','arguments':json.dumps({'query':'项目 A','as_of':''})}}]}
        with patch('assistant.runtime.chat',side_effect=[(tool,{}),({'role':'assistant','content':'图像采集采用 ESP32-P4，负责人是硬件组张工。'}, {})]) as chat:
            run_turn(t2.pk)
            self.assertIn('你想查哪个项目',json.dumps(chat.call_args_list[0].args[0],ensure_ascii=False))
            self.assertIn('ESP32-P4',json.dumps(chat.call_args_list[1].args[0],ensure_ascii=False))
        t2.refresh_from_db();self.assertEqual(t2.state,'done')
        self.assertEqual(c.messages.count(),4)

    def test_proposal_is_private_until_explicit_confirmation(self):
        p=propose(self.alice,card())
        self.assertFalse(EnterpriseProject.objects.exists())
        self.client.force_authenticate(self.bob)
        self.assertEqual(self.client.get('/api/assistant/proposals/').data,[])
        self.assertEqual(self.client.post(f'/api/assistant/proposals/{p.pk}/',{'action':'confirm'},format='json').status_code,404)
        self.client.force_authenticate(self.alice)
        self.assertEqual(self.client.post(f'/api/assistant/proposals/{p.pk}/',{},format='json').status_code,400)
        self.assertEqual(self.client.post(f'/api/assistant/proposals/{p.pk}/',{'action':'confirm'},format='json').status_code,200)
        confirm_proposal(p.pk,self.alice)
        self.assertEqual(ProjectVersion.objects.count(),1)
        self.client.force_authenticate(self.bob)
        self.assertEqual(len(self.client.get('/api/assistant/projects/').data['items']),1)

    def test_cross_owner_edit_needs_admin_and_conflict_does_not_overwrite(self):
        project=self.create_project()
        p=propose(self.bob,card(chip='STM32H7'),project.pk,1)
        result=confirm_proposal(p.pk,self.bob)
        self.assertEqual(result.state,'needs_admin')
        self.assertEqual(project_record(project)['data']['modules'][0]['chips'],['ESP32-P4'])
        own=propose(self.alice,card(chip='RK3568'),project.pk,1)
        confirm_proposal(own.pk,self.alice)
        with self.assertRaises(ValueError):confirm_proposal(p.pk,self.admin)
        project.refresh_from_db();self.assertEqual(project.revision,2)
        fresh=propose(self.bob,card(chip='STM32H7'),project.pk,2)
        confirm_proposal(fresh.pk,self.bob);confirm_proposal(fresh.pk,self.admin)
        project.refresh_from_db();self.assertEqual(project.revision,3)

    def test_historical_effective_dates_and_alias(self):
        project=self.create_project()
        p=propose(self.alice,card(chip='RK3568',effective='2099-01-01'),project.pk,1)
        confirm_proposal(p.pk,self.alice);project.refresh_from_db()
        self.assertEqual(project_record(project,'2026-10-01')['data']['modules'][0]['chips'],['ESP32-P4'])
        self.assertEqual(project_record(project,'2099-02-01')['data']['modules'][0]['chips'],['RK3568'])
        self.assertEqual(self.client.get('/api/assistant/projects/?q=视觉设备').data['items'][0]['id'],project.pk)
        self.assertEqual(self.client.get('/api/assistant/projects/?as_of=bad').status_code,400)

    def test_document_preview_confirmation_and_withdrawal(self):
        d,v=self.document();self.assertEqual(v.state,'ready');self.assertEqual(retrieve('图像采集'),[])
        self.client.force_authenticate(self.bob)
        self.assertEqual(self.client.get(f'/api/assistant/documents/{d.pk}/').status_code,404)
        self.assertEqual(self.client.get(f'/api/assistant/documents/{d.pk}/file/').status_code,404)
        with self.assertRaises(PermissionError):confirm_document(v.pk,self.bob)
        confirm_document(v.pk,self.alice)
        self.assertEqual(retrieve('图像采集')[0]['version_id'],v.pk)
        self.assertEqual(self.client.get(f'/api/assistant/documents/{d.pk}/file/').status_code,200)
        self.client.force_authenticate(self.alice)
        self.client.post(f'/api/assistant/documents/{d.pk}/',{'action':'withdraw','version_id':v.pk},format='json')
        self.assertEqual(retrieve('图像采集'),[])
        self.client.force_authenticate(self.bob)
        self.assertEqual(self.client.get(f'/api/assistant/documents/{d.pk}/file/').status_code,404)

    def test_document_revisions_keep_sources(self):
        d,v=self.document();confirm_document(v.pk,self.alice)
        newer=KnowledgeVersion.objects.create(document=d,revision=2,state='indexing',text='项目 A 图像采集改用 RK3568。')
        index_document(newer.pk);confirm_document(newer.pk,self.alice)
        v.refresh_from_db();self.assertEqual(v.state,'archived')
        self.assertEqual({x['version_id'] for x in retrieve('图像采集')},{newer.pk})
        self.client.force_authenticate(self.bob)
        self.assertEqual(self.client.get(f'/api/assistant/documents/{d.pk}/file/?version={v.pk}').status_code,200)

    def test_source_document_must_be_published(self):
        d,v=self.document();data=card();data['document_ids']=[d.pk]
        r=self.client.post('/api/assistant/proposals/',{'data':data},format='json')
        self.assertEqual(r.status_code,400)
        confirm_document(v.pk,self.alice)
        self.assertEqual(self.client.post('/api/assistant/proposals/',{'data':data},format='json').status_code,201)

    def test_project_search_includes_confirmed_card_documents_only(self):
        project=self.create_project()
        d,v=self.document();confirm_document(v.pk,self.alice)
        data=card();data['document_ids']=[d.pk]
        proposal=propose(self.alice,data,project.pk,project.revision)
        self.assertEqual(retrieve('图像采集',project.pk),[])
        confirm_proposal(proposal.pk,self.alice)
        self.assertEqual({x['document_id'] for x in retrieve('图像采集',project.pk)},{d.pk})
        d.refresh_from_db();self.assertIsNone(d.project_id)
        self.client.post(f'/api/assistant/documents/{d.pk}/',{'action':'withdraw','version_id':v.pk},format='json')
        self.assertEqual(retrieve('图像采集',project.pk),[])

    def test_project_search_uses_effective_card_and_direct_upload_links(self):
        project=self.create_project()
        direct,direct_v=self.document();direct.project=project;direct.save();confirm_document(direct_v.pk,self.alice)
        linked,linked_v=self.document();confirm_document(linked_v.pk,self.alice)
        future,future_v=self.document();confirm_document(future_v.pk,self.alice)
        data=card();data['document_ids']=[linked.pk]
        confirm_proposal(propose(self.alice,data,project.pk,1).pk,self.alice)
        data=card(effective='2099-01-01');data['document_ids']=[future.pk]
        confirm_proposal(propose(self.alice,data,project.pk,2).pk,self.alice)
        self.assertEqual({x['document_id'] for x in retrieve('图像采集',project.pk)},{direct.pk,linked.pk})
        self.assertEqual(retrieve('图像采集',project.pk+1000),[])

    def test_docx_xlsx_csv_and_markdown(self):
        from docx import Document
        from openpyxl import Workbook
        word=Document();word.add_paragraph('采集模块由张工负责');buf=io.BytesIO();word.save(buf)
        book=Workbook();book.active.append(['模块','负责人']);book.active.append(['采集','张工']);excel=io.BytesIO();book.save(excel)
        fixtures=[('a.docx',buf.getvalue()),('a.xlsx',excel.getvalue()),('a.csv','模块,负责人\n采集,张工'.encode()),('a.md','采集模块由张工负责'.encode())]
        for name,content in fixtures:
            response=self.client.post('/api/assistant/documents/',{'title':name,'file':SimpleUploadedFile(name,content)},format='multipart')
            self.assertEqual(response.status_code,202,response.data)
            v=KnowledgeVersion.objects.get(pk=response.data['versions'][0]['id']);v.state='indexing';v.save();index_document(v.pk);v.refresh_from_db()
            self.assertEqual(v.state,'ready',v.error);self.assertIn('张工',v.text)

    def test_public_query_boundary_and_ssrf(self):
        self.create_project()
        with patch('assistant.tools.httpx.Client') as client:
            for terms,model in [(['项目 A'],'ESP32'),(['MCU'],'PROJ-A'),(['MCU'],'张工'),(['MCU'],'hello world')]:
                with self.assertRaises(ValueError):search_chips(terms,model)
            client.assert_not_called()
        for url in ['http://example.com','https://user:pass@example.com','https://example.com:8443','file:///etc/passwd']:
            with self.assertRaises(ValueError):public_url(url)
        t=Turn.objects.create(conversation=Conversation.objects.create(owner=self.alice),state='running')
        with self.assertRaises(ValueError):execute('read_source',{'url':'https://127.0.0.1'},t,[])

    def test_batch_research_collects_bounded_sources_and_uses_source_ids(self):
        c=Conversation.objects.create(owner=self.alice);t=Turn.objects.create(conversation=c,state='running')
        fixture={'items':[{'url':'https://example.test/chip','title':'公开资料','kind':'search_excerpt','text':'摘要'}],
                 'evidence':[{'url':'https://example.test/chip','title':'公开资料','kind':'source_text','text':'已读取正文','notice':''}],'warnings':[]}
        sources=[]
        with patch('assistant.tools._research_query',side_effect=lambda _:__import__('copy').deepcopy(fixture)):
            result=execute('research_chips',{'models':['ESP32-P4','STM32N6'],'terms':['MCU','MIPI CSI-2']},t,sources)
        self.assertEqual(len(result['candidates']),2)
        self.assertEqual([x['source_id'] for x in sources],[1,2,3,4])
        with patch('assistant.tools.read_source',return_value=dict(fixture['evidence'][0])) as read:
            execute('read_source',{'source_id':1},t,sources)
            read.assert_called_once_with('https://example.test/chip')
        with self.assertRaises(ValueError):execute('read_source',{'source_id':999},t,sources)
        self.create_project()
        with patch('assistant.tools._research_pool.submit') as submit:
            with self.assertRaises(ValueError):execute('research_chips',{'models':['ESP32-P4','PROJ-A'],'terms':['MCU']},t,[])
            submit.assert_not_called()

    def test_requirements_corrections_persist_privately_across_turns(self):
        c=Conversation.objects.create(owner=self.alice,brief='离线视觉设备；成本未定')
        t=Turn.objects.create(conversation=c,state='running')
        execute('remember_requirements',{'brief':'离线视觉设备；用户将成本上限修正为 200 元'},t,[])
        c.refresh_from_db();self.assertIn('200 元',c.brief)
        Message.objects.create(conversation=c,role='user',content='按修正后的成本继续')
        with patch('assistant.runtime.chat',return_value=({'role':'assistant','content':'已使用修正后的成本条件'},{})) as chat:
            run_turn(t.pk)
            self.assertIn('200 元',chat.call_args.args[0][0]['content'])
        self.assertFalse(EnterpriseProject.objects.exists())
        self.assertFalse(KnowledgeDocument.objects.exists())
        self.client.force_authenticate(self.bob)
        self.assertEqual(self.client.get(f'/api/assistant/conversations/{c.pk}/').status_code,404)

    @override_settings(AGENT_MAX_STEPS=3)
    def test_completed_batch_delivers_first_result_without_search_expansion(self):
        c=Conversation.objects.create(owner=self.alice)
        Message.objects.create(conversation=c,role='user',content='比较三个芯片，缺失参数注明未知')
        t=Turn.objects.create(conversation=c,state='running')
        calls=[]
        def model(messages,allow_tools=True,timeout=None):
            calls.append(allow_tools)
            if not allow_tools:
                return {'role':'assistant','content':'已核实两款，另一款资料无法读取，性能待实测。'},{}
            name='research_chips' if len(calls)==1 else 'search_chips'
            return {'role':'assistant','content':None,'tool_calls':[{'id':str(len(calls)),'type':'function','function':{'name':name,'arguments':'{}'}}]},{}
        with patch('assistant.runtime.chat',side_effect=model), patch('assistant.runtime.execute',return_value={'candidates':[{'model':'ESP32-P4','evidence':['source']},{'model':'STM32N6','evidence':[]}]}) as execute:
            run_turn(t.pk)
        t.refresh_from_db()
        self.assertEqual(calls,[True,False])
        self.assertEqual(execute.call_count,1)
        self.assertEqual(t.state,'done')
        self.assertIn('待实测',c.messages.last().content)

    def test_correction_is_retained_when_model_claims_saved_but_skips_tool(self):
        c=Conversation.objects.create(owner=self.alice,brief='离线视觉设备，5 寸屏，预算 200 元')
        correction='修正条件：不要 Linux，只考虑 RTOS 或裸机。请更新需求摘要。'
        Message.objects.create(conversation=c,role='user',content=correction)
        t=Turn.objects.create(conversation=c,state='running')
        with patch('assistant.runtime.chat',return_value=({'role':'assistant','content':'已更新私有需求摘要。'},{})):
            run_turn(t.pk)
        c.refresh_from_db()
        self.assertIn(correction,c.brief)
        self.assertIn('预算 200 元',c.brief)
        self.assertIn('原文',c.brief)
        Message.objects.create(conversation=c,role='user',content='沿用最新条件继续')
        t2=Turn.objects.create(conversation=c,state='running')
        with patch('assistant.runtime.chat',return_value=({'role':'assistant','content':'继续 RTOS 路线。'},{})) as chat:
            run_turn(t2.pk)
            self.assertIn(correction,chat.call_args.args[0][0]['content'])
        self.assertFalse(EnterpriseProject.objects.exists())
        self.client.force_authenticate(self.bob)
        self.assertEqual(self.client.get(f'/api/assistant/conversations/{c.pk}/').status_code,404)

    def test_successful_summary_tool_keeps_its_consolidated_summary(self):
        c=Conversation.objects.create(owner=self.alice,brief='旧预算 300 元')
        Message.objects.create(conversation=c,role='user',content='预算改为 200 元')
        t=Turn.objects.create(conversation=c,state='running')
        tool={'role':'assistant','content':None,'tool_calls':[{'id':'memory','type':'function','function':{'name':'remember_requirements','arguments':json.dumps({'brief':'用户确认：预算 200 元'})}}]}
        with patch('assistant.runtime.chat',side_effect=[(tool,{}),({'role':'assistant','content':'已改为 200 元。'}, {})]):
            run_turn(t.pk)
        c.refresh_from_db();self.assertEqual(c.brief,'用户确认：预算 200 元')

    @override_settings(AGENT_RUNTIME='codex',LLM_BASE_URL='https://provider.example/v1',LLM_API_KEY='test-only',LLM_MODEL='test')
    def test_codex_uses_same_private_write_boundary(self):
        c=Conversation.objects.create(owner=self.alice);Message.objects.create(conversation=c,role='user',content='保存项目资料')
        t=Turn.objects.create(conversation=c,state='running')
        def codex(**kwargs):
            result=kwargs['execute_tool']('propose_project',{'project_id':0,'base_revision':0,'data':json.dumps(card())})
            self.assertIn('proposal_id',result);self.assertFalse(EnterpriseProject.objects.exists())
            return '已整理待确认内容，请确认。',{'input_tokens':123,'output_tokens':20}
        with patch('assistant.app_server.run_codex',side_effect=codex):run_turn(t.pk)
        t.refresh_from_db();self.assertEqual(t.state,'done');self.assertEqual(c.proposals.count(),1)
